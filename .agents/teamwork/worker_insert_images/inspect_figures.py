import docx
import sys

# Force UTF-8 stdout
sys.stdout.reconfigure(encoding='utf-8')

for doc_path in [r'c:\IT_System\BaoCao_ITTMS_TKPMHDT (4).docx', r'c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx']:
    doc = docx.Document(doc_path)
    print(f"\n==========================================")
    print(f"FILE: {doc_path}")
    print(f"Total paragraphs: {len(doc.paragraphs)}")
    headings = [p for p in doc.paragraphs if p.style.name.startswith('Heading')]
    print(f"Total headings: {len(headings)}")
    print(f"Total tables: {len(doc.tables)}")
    
    # List all figures
    figure_count = 0
    for idx, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith('Hình ') and ('–' in txt or '-' in txt or '.' in txt):
            figure_count += 1
            prev_p = doc.paragraphs[idx - 1] if idx > 0 else None
            prev_style = prev_p.style.name if prev_p else 'None'
            prev_xml = prev_p._element.xml if prev_p else ''
            prev_has_drawing = '<w:drawing' in prev_xml
            curr_has_drawing = '<w:drawing' in p._element.xml
            print(f"{figure_count:02d}. [P{idx:03d}] {txt}")
            print(f"    Prev: [P{idx-1:03d}] style={prev_style:<16} has_drawing={prev_has_drawing} text_len={len(prev_p.text)}")
            if curr_has_drawing:
                print(f"    Curr: HAS DRAWING IN CAPTION ITSELF!")
    print(f"Total detected figure captions: {figure_count}")
