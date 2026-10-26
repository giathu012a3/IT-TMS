import os
import sys
import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

sys.stdout.reconfigure(encoding='utf-8')

DOCS_TO_UPDATE = [
    r"c:\IT_System\BaoCao_ITTMS_TKPMHDT (4).docx",
    r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx"
]

DIAGRAM_DIR = r"c:\IT_System\docs\diagrams"

# 6 missing: insert before caption
MISSING_DIAGRAMS = {
    "Hình 3.1. Lược đồ thực thể – quan hệ (ERD) của IT-TMS": ("Hinh3_1_ERD_8Bang.png", 6.2),
    "Hình 5.1. Quy trình xử lý ticket từ lúc tạo đến khi kết thúc": ("Hinh5_1_BPMN_TicketLifecycle.png", 6.0),
    "Hình 5.2. Sơ đồ chuyển trạng thái của ticket": ("Hinh5_2_StateMachine_Ticket.png", 6.0),
    "Hình 5.3. Lược đồ lớp tổng quát thể hiện các mẫu thiết kế": ("Hinh5_3_ClassDiagram_MVC.png", 6.2),
    "Hình 5.6. Lược đồ tuần tự – UC09 – Phân công yêu cầu": ("Hinh5_6_Sequence_UC09_PhanCong.png", 6.0),
    "Hình 6.1. Cấu trúc State Pattern cho ticket": ("Hinh6_1_StatePattern_ClassDiagram.png", 6.0),
}

# 13 outdated: replace drawing in prev paragraph
OUTDATED_DIAGRAMS = {
    "Hình 2.1. Lược đồ use case chính của hệ thống IT-TMS": ("Hinh2_1_UseCaseTongThe.png", 6.0),
    "Hình 2.2. Use case chi tiết – Actor User": ("Hinh2_2_UseCaseUser.png", 6.0),
    "Hình 2.3. Use case chi tiết – Actor Staff": ("Hinh2_4_UseCaseStaff.png", 6.0),
    "Hình 2.4. Use case chi tiết – Actor Leader": ("Hinh2_3_UseCaseLeader.png", 6.0),
    "Hình 2.5. Use case chi tiết – Actor Admin": ("Hinh2_5_UseCaseAdmin.png", 6.0),
    "Hình 5.4. Lược đồ tuần tự – UC01 – Đăng nhập": ("Hinh5_4_Sequence_UC01_Login.png", 6.0),
    "Hình 5.5. Lược đồ tuần tự – UC03 – Tạo yêu cầu mới": ("Hinh5_5_Sequence_UC03_CreateTicket.png", 6.0),
    "Hình 5.7. Lược đồ tuần tự – UC07 – Cập nhật trạng thái xử lý": ("Hinh5_7_Sequence_UC07_ProcessTicket.png", 6.0),
    "Hình 5.8. Lược đồ tuần tự – UC06 – Trao đổi thông tin ticket": ("Hinh5_8_Sequence_UC06_ChatRealtime.png", 6.0),
    "Hình 5.9. Lược đồ tuần tự – UC10 – Từ chối yêu cầu": ("Hinh5_9_Sequence_UC10_RejectTicket.png", 6.0),
    "Hình 5.10. Lược đồ tuần tự – UC19, UC04 – Đóng / Mở lại yêu cầu và đánh giá": ("Hinh5_10_Sequence_UC04_19_Feedback.png", 6.0),
    "Hình 5.11. Lược đồ tuần tự – UC15 – Kích hoạt / Vô hiệu hóa tài khoản": ("Hinh5_11_Sequence_UC15_ToggleUser.png", 6.0),
    "Hình 5.12. Lược đồ tuần tự – UC11 – Xem tổng quan hoạt động (Dashboard Leader)": ("Hinh5_12_Sequence_UC11_DashboardLeader.png", 6.0),
}

def process_document(doc_path):
    print(f"\n=======================================================")
    print(f"XỬ LÝ CHÈN VÀ THAY THẾ ẢNH: {doc_path}")
    print(f"=======================================================")
    doc = docx.Document(doc_path)
    
    replaced_count = 0
    inserted_count = 0
    
    for p in list(doc.paragraphs):
        txt = p.text.strip()
        
        # Check outdated diagrams
        if txt in OUTDATED_DIAGRAMS:
            img_name, width_inch = OUTDATED_DIAGRAMS[txt]
            img_path = os.path.join(DIAGRAM_DIR, img_name)
            
            # Get preceding paragraph element
            prev_elem = p._element.getprevious()
            assert prev_elem is not None, f"No previous element for {txt}"
            assert '<w:drawing' in prev_elem.xml, f"Prev para for {txt} has no drawing!"
            
            prev_p = docx.text.paragraph.Paragraph(prev_elem, doc)
            orig_style = prev_p.style.name
            
            # Clear existing runs
            for r in list(prev_elem.xpath('./w:r')):
                prev_elem.remove(r)
            
            # Add new picture run
            new_run = prev_p.add_run()
            new_run.add_picture(img_path, width=Inches(width_inch))
            prev_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            prev_p.style = orig_style
            
            # Ensure caption is centered
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            replaced_count += 1
            print(f"[REPLACED] {txt[:40]}... -> {img_name}")
            
        elif txt in MISSING_DIAGRAMS:
            img_name, width_inch = MISSING_DIAGRAMS[txt]
            img_path = os.path.join(DIAGRAM_DIR, img_name)
            
            # Insert new paragraph before caption
            new_p = p.insert_paragraph_before()
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_p.style = doc.styles['Normal']
            
            run = new_p.add_run()
            run.add_picture(img_path, width=Inches(width_inch))
            
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            inserted_count += 1
            print(f"[INSERTED] {txt[:40]}... -> {img_name}")
            
    assert replaced_count == 13, f"Expected 13 replacements, got {replaced_count}"
    assert inserted_count == 6, f"Expected 6 insertions, got {inserted_count}"
    print(f"\nThành công: Đã thay thế {replaced_count}/13 ảnh cũ và chèn mới {inserted_count}/6 ảnh thiếu!")
    
    doc.save(doc_path)
    print(f"Đã lưu tệp Word cập nhật: {doc_path}")

if __name__ == "__main__":
    for doc_path in DOCS_TO_UPDATE:
        process_document(doc_path)
    print("\n>>> HOÀN TẤT CẬP NHẬT CẢ 2 TỆP TIN WORD CHÍNH THỨC! <<<")
