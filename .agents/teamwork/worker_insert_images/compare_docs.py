import docx
import sys
sys.stdout.reconfigure(encoding='utf-8')

doc1 = docx.Document(r'c:\IT_System\BaoCao_ITTMS_TKPMHDT (4).docx')
doc2 = docx.Document(r'c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx')

print(f"Doc1 paras: {len(doc1.paragraphs)}, Doc2 paras: {len(doc2.paragraphs)}")
h1 = [p for p in doc1.paragraphs if p.style.name.startswith("Heading")]
h2 = [p for p in doc2.paragraphs if p.style.name.startswith("Heading")]
print(f"Doc1 headings: {len(h1)}, Doc2 headings: {len(h2)}")

diffs = []
for i in range(min(len(doc1.paragraphs), len(doc2.paragraphs))):
    p1 = doc1.paragraphs[i]
    p2 = doc2.paragraphs[i]
    if p1.style.name != p2.style.name:
        diffs.append((i, f"style: {p1.style.name} vs {p2.style.name}"))
    if p1.text.strip() != p2.text.strip():
        if p1.text[:20] != p2.text[:20]:
            diffs.append((i, f"text: {repr(p1.text[:30])} vs {repr(p2.text[:30])}"))

print(f"Total diffs in paragraphs: {len(diffs)}")
for d in diffs[:10]:
    print("  Diff:", d)
