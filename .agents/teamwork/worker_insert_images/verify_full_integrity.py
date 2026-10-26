import os
import sys
import zipfile
import re
import xml.etree.ElementTree as ET
import docx

sys.stdout.reconfigure(encoding='utf-8')

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
ns = {"w": W_NS}

DOCS_TO_VERIFY = [
    ("Doc1", r"c:\IT_System\BaoCao_ITTMS_TKPMHDT (4).docx"),
    ("Doc2", r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx")
]

OFF_LIMIT_WORDS = ["unit test", "unit-test", "devops", "kiểm thử tự động"]

def audit_document(label, file_path):
    print(f"\n=======================================================")
    print(f"KIỂM TOÁN TÍNH TOÀN VẸN: {label} ({file_path})")
    print(f"=======================================================")
    
    assert os.path.exists(file_path), f"File {file_path} không tồn tại!"
    file_size = os.path.getsize(file_path)
    print(f"1. Dung lượng file: {file_size:,} bytes ({file_size/1024/1024:.2f} MB)")
    
    # 2. Check XML synthetic namespaces & drawings count
    synthetic_matches = []
    with zipfile.ZipFile(file_path) as z:
        for name in z.namelist():
            if name.endswith(".xml") or name.endswith(".rels"):
                content = z.read(name).decode("utf-8", errors="ignore")
                matches = re.findall(r"(ns\d+:|xmlns:ns\d+)", content)
                if matches:
                    synthetic_matches.append((name, set(matches)))
        xml_data = z.read("word/document.xml")
        root = ET.fromstring(xml_data)
        
        drawing_count = len(root.findall(".//w:drawing", ns))
        print(f"2. Số thẻ <w:drawing> trong document.xml: {drawing_count} (Kỳ vọng: 41)")
        assert drawing_count == 41, f"LỖI: Kỳ vọng 41 drawings, phát hiện {drawing_count}!"
        
    print(f"   Tiền tố namespace nhân tạo: {synthetic_matches} (count={len(synthetic_matches)})")
    assert len(synthetic_matches) == 0, "LỖI: Phát hiện synthetic namespaces lạ!"

    # 3. Headings và TOC
    doc = docx.Document(file_path)
    paras = doc.paragraphs
    headings = [p for p in paras if p.style.name.startswith("Heading")]
    h1 = [p for p in headings if p.style.name == 'Heading 1']
    h2 = [p for p in headings if p.style.name == 'Heading 2']
    h3 = [p for p in headings if p.style.name == 'Heading 3']
    text_headings = [p for p in headings if p.text.strip()]
    draw_headings = [p for p in headings if not p.text.strip()]
    
    print(f"3. CẤU TRÚC ĐOẠN VĂN & TIÊU ĐỀ:")
    print(f"   - Tổng số Paragraphs: {len(paras)} (Kỳ vọng: 775)")
    assert len(paras) == 775, f"LỖI: Paragraphs {len(paras)} != 775"
    
    print(f"   - Tổng số Headings: {len(headings)} (Kỳ vọng: 149)")
    print(f"     + Heading 1: {len(h1)} (Kỳ vọng: 14)")
    print(f"     + Heading 2: {len(h2)} (Kỳ vọng: 62)")
    print(f"     + Heading 3: {len(h3)} (Kỳ vọng: 73)")
    print(f"     + Tiêu đề có text: {len(text_headings)} (Kỳ vọng: 136)")
    print(f"     + Heading rỗng (Image holder): {len(draw_headings)} (Kỳ vọng: 13)")
    assert len(headings) == 149, f"LỖI: Headings {len(headings)} != 149"
    assert len(text_headings) == 136, f"LỖI: Text Headings {len(text_headings)} != 136"
    assert len(draw_headings) == 13, f"LỖI: Draw Headings {len(draw_headings)} != 13"
    
    # 4. Check TOC items (paragraphs 60 to 196)
    toc_paras = [p.text.strip() for p in paras[60:196] if p.text.strip()]
    print(f"4. CẤU TRÚC MỤC LỤC (TOC):")
    print(f"   - Số mục TOC (paras 60-196): {len(toc_paras)} (Kỳ vọng: 136)")
    assert len(toc_paras) == 136, f"LỖI: TOC items {len(toc_paras)} != 136"
    
    # 5. Check tables
    print(f"5. CẤU TRÚC BẢNG BIỂU:")
    print(f"   - Tổng số bảng: {len(doc.tables)} (Kỳ vọng: 55)")
    assert len(doc.tables) == 55, f"LỖI: Số bảng {len(doc.tables)} != 55"
    
    # 6. Kiểm tra Bảng PL.1
    pl1 = doc.tables[-1]
    print(f"6. KIỂM TRA BẢNG PHÂN CÔNG PHỤ LỤC 1 (BẢNG PL.1):")
    assert len(pl1.rows) == 5, f"LỖI: Số dòng {len(pl1.rows)} != 5"
    assert len(pl1.columns) == 3, f"LỖI: Số cột {len(pl1.columns)} != 3"
    
    expected_members = [
        ("Nguyễn Gia Thụ", "25810041", "Staff", "StaffController.cs", "State Pattern", "5.2", "100%"),
        ("Nguyễn Mạnh Tiến", "25810043", "Leader", "LeaderController.cs", "Factory", "5.1", "100%"),
        ("Hồ Thanh Tình", "25810044", "User", "TicketController.cs", "Observer", "TicketHub", "100%"),
        ("Trần Thanh Danh", "25810009", "Admin", "AdminController.cs", "Decorator", "ERD", "100%")
    ]
    all_table_text = ""
    for r_idx, exp in enumerate(expected_members, start=1):
        row = pl1.rows[r_idx]
        name = row.cells[0].text.strip()
        mssv = row.cells[1].text.strip()
        work = row.cells[2].text.strip()
        all_table_text += " " + name + " " + mssv + " " + work
        assert exp[0] in name, f"LỖI: Tên {name} không chứa {exp[0]}"
        assert exp[1] in mssv, f"LỖI: MSSV {mssv} không chứa {exp[1]}"
        for kw in exp[2:]:
            assert kw.lower() in work.lower(), f"LỖI: Dòng của {name} thiếu từ khóa {kw}"
            
    for bad_word in OFF_LIMIT_WORDS:
        assert bad_word not in all_table_text.lower(), f"LỖI: Bảng PL.1 còn chứa từ khóa cấm: '{bad_word}'"
    print("   - Bảng PL.1: 4 Vertical Slices độc lập, 100% tỷ lệ đóng góp, 0 từ khóa ngoài lề!")
    
    # 7. Check All 41 Figures in body
    print(f"7. XÁC MINH DANH MỤC 41 HÌNH ẢNH TRONG NỘI DUNG:")
    figure_found = 0
    for idx, p in enumerate(paras):
        txt = p.text.strip()
        if txt.startswith("Hình ") and ("–" in txt or "-" in txt or "." in txt):
            prev_elem = p._element.getprevious()
            if prev_elem is not None and '<w:drawing' in prev_elem.xml:
                figure_found += 1
                fig_num = txt.split('.')[0] if '.' in txt else txt[:8]
                # Spot check print
                if figure_found in [1, 5, 6, 7, 28, 29, 30, 31, 32, 34, 40, 41]:
                    print(f"   + [Fig {figure_found:02d}] P{idx:03d}: {txt[:60]}... -> OK")
                
    print(f"   - Tổng số hình ảnh có node <w:drawing> ngay phía trên: {figure_found}/41")
    assert figure_found == 41, f"LỖI: Kỳ vọng 41 hình ảnh, đếm được {figure_found}"
    print(f"[VERDICT] {label}: ĐẠT CHUẨN TOÀN DIỆN 100% (PASS)")

if __name__ == "__main__":
    for label, path in DOCS_TO_VERIFY:
        audit_document(label, path)
    print("\n>>> TẤT CẢ CÁC BÀI KIỂM TOÁN TÍNH TOÀN VẸN ĐỀU VƯỢT QUA 100% TUYỆT ĐỐI! <<<")
