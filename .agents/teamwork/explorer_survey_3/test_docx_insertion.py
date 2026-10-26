import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont
import shutil
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

test_dir = r"c:\IT_System\.agents\teamwork\explorer_survey_3"
source_docx = r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx"
test_docx = os.path.join(test_dir, "test_sample.docx")
test_img = os.path.join(test_dir, "dummy_test.png")

# 1. Create a dummy test image with PIL at 300 DPI
img = Image.new('RGB', (1800, 1000), color=(245, 247, 250))
draw = ImageDraw.Draw(img)
draw.rectangle([20, 20, 1780, 980], outline=(37, 99, 235), width=6)
draw.text((100, 100), "TEST DIAGRAM (300 DPI EQUIVALENT)", fill=(15, 23, 42))
img.save(test_img, 'PNG', dpi=(300, 300))
print(f"Created dummy image: {test_img} (1800x1000, 300 DPI)")

# 2. Copy source docx
shutil.copyfile(source_docx, test_docx)
print(f"Copied source docx to: {test_docx}")

# 3. Test python-docx modification
doc = docx.Document(test_docx)

# Test Case B: Insert image before Hình 3.1
target_caption = "Hình 3.1. Lược đồ thực thể – quan hệ (ERD) của IT-TMS"
found = False
for i, p in enumerate(doc.paragraphs):
    if p.text.strip() == target_caption:
        print(f"Found target caption at P#{i}: {p.text}")
        img_p = p.insert_paragraph_before()
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        # Ensure style is Normal
        img_p.style = doc.styles['Normal']
        run = img_p.add_run()
        run.add_picture(test_img, width=Inches(6.0))
        found = True
        break

if not found:
    print("ERROR: Target caption not found!")

# Save modified docx
output_docx = os.path.join(test_dir, "test_modified.docx")
doc.save(output_docx)
print(f"Saved modified docx to: {output_docx}")

# 4. Verify with Word COM
import subprocess
ps_script = f"""
$w = New-Object -ComObject Word.Application
$w.Visible = $false
try {{
    $doc = $w.Documents.Open('{output_docx}', $false, $true)
    Write-Host 'Word COM Opened Successfully!'
    Write-Host 'Paragraph count:' $doc.Paragraphs.Count
    Write-Host 'InlineShapes count:' $doc.InlineShapes.Count
    $doc.Close($false)
    Write-Host 'VERIFICATION PASSED: 0 XML CORRUPTION!'
}} catch {{
    Write-Host 'VERIFICATION FAILED:' $_.Exception.Message
}} finally {{
    $w.Quit()
}}
"""
ps_file = os.path.join(test_dir, "verify_word.ps1")
with open(ps_file, "w", encoding="utf-8") as f:
    f.write(ps_script)

res = subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", ps_file], capture_output=True, text=True)
print("Word COM Output:\n" + res.stdout)
if res.stderr:
    print("Word COM Stderr:\n" + res.stderr)
