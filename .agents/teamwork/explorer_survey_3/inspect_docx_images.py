import docx
import zipfile
import os
import sys
import json
import re

sys.stdout.reconfigure(encoding='utf-8')

def extract_doc_structure(file_path):
    print("=" * 70)
    print(f"ANALYZING: {os.path.basename(file_path)}")
    print("=" * 70)
    
    doc = docx.Document(file_path)
    
    # Extract docx relationships to map rId -> target media
    rels = {}
    for rId, rel in doc.part.rels.items():
        if "image" in rel.target_ref:
            rels[rId] = rel.target_ref
            
    print(f"Total image relationships: {len(rels)}")
    
    # Analyze paragraphs
    figures = []
    # Regular expression for figure captions: e.g. "Hình 2.1:", "Hình 3.1", "Hình 5.1", etc.
    fig_pattern = re.compile(r'^(Hình\s+\d+\.\d+.*)$', re.IGNORECASE)
    
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        xml = p._element.xml
        
        # Check drawings in this paragraph
        drawings = []
        if '<w:drawing' in xml:
            # find blip r:embed
            matches = re.findall(r'r:embed="([^"]+)"', xml)
            drawings.extend(matches)
            
        m = fig_pattern.match(text)
        if m:
            caption_text = m.group(1)
            # check previous paragraph for drawing
            prev_p = doc.paragraphs[i-1] if i > 0 else None
            prev_xml = prev_p._element.xml if prev_p else ""
            prev_drawings = re.findall(r'r:embed="([^"]+)"', prev_xml) if prev_p else []
            
            figures.append({
                "index": i,
                "caption": caption_text,
                "drawings_in_caption": drawings,
                "drawings_in_prev_p": prev_drawings,
                "prev_p_text": prev_p.text.strip() if prev_p else "",
                "has_drawing": bool(drawings or prev_drawings),
                "media_targets": [rels.get(r, "UNKNOWN") for r in (drawings or prev_drawings)]
            })
            
    print(f"Found {len(figures)} figure captions matching 'Hình X.Y':")
    for f in figures:
        status = "OK" if f["has_drawing"] else "MISSING"
        targets = ", ".join(f["media_targets"]) if f["media_targets"] else "NONE"
        print(f"  P#{f['index']:3d} [{status:7s}] {f['caption'][:50]:<50} -> {targets}")
        
    return figures

if __name__ == '__main__':
    doc1 = r"c:\IT_System\BaoCao_ITTMS_TKPMHDT (4).docx"
    doc2 = r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx"
    figs1 = extract_doc_structure(doc1)
    figs2 = extract_doc_structure(doc2)
