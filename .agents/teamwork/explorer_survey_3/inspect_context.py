import docx
import sys

sys.stdout.reconfigure(encoding='utf-8')

def inspect_context(docx_path, targets):
    print("=" * 70)
    print(f"CONTEXT INSPECTION: {docx_path}")
    print("=" * 70)
    doc = docx.Document(docx_path)
    for t in targets:
        for i, p in enumerate(doc.paragraphs):
            if t in p.text and ('Chương' not in p.text and 'Danh mục' not in p.text and len(p.text) < 150):
                print(f"\n--- TARGET: {t} at P#{i} ---")
                start = max(0, i - 4)
                end = min(len(doc.paragraphs), i + 4)
                for j in range(start, end):
                    marker = ">>>>> " if j == i else "      "
                    xml = doc.paragraphs[j]._element.xml
                    has_dw = "[HAS_DRAWING]" if "<w:drawing" in xml else "             "
                    style = doc.paragraphs[j].style.name
                    print(f"{marker}P#{j:3d} ({style:<15}) {has_dw} : {doc.paragraphs[j].text.strip()[:90]}")

if __name__ == '__main__':
    targets = ["Hình 3.1", "Hình 5.1", "Hình 5.2", "Hình 5.3", "Hình 5.6", "Hình 6.1"]
    inspect_context(r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx", targets)
