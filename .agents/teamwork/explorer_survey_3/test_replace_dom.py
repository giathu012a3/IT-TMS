import docx
from docx.shared import Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw
import shutil
import os
import sys

sys.stdout.reconfigure(encoding='utf-8')

test_dir = r"c:\IT_System\.agents\teamwork\explorer_survey_3"
source_docx = r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx"
test_docx = os.path.join(test_dir, "test_replace.docx")
test_img = os.path.join(test_dir, "dummy_replace.png")

# Create a test image
img = Image.new('RGB', (1800, 1200), color=(255, 255, 255))
draw = ImageDraw.Draw(img)
draw.rectangle([10, 10, 1790, 1190], outline=(16, 185, 129), width=5)
draw.text((100, 100), "REPLACED DIAGRAM TEST", fill=(0, 0, 0))
img.save(test_img, 'PNG', dpi=(300, 300))

shutil.copyfile(source_docx, test_docx)
doc = docx.Document(test_docx)

# P#357 is the image paragraph for Hình 2.1
p = doc.paragraphs[357]
print("Before replacement, p XML has drawing:", "<w:drawing" in p._element.xml)

# Replace drawing: clear element children except pPr or clear runs
# Better: remove all child r elements or clear text
# In oxml:
for r in list(p._element.xpath('./w:r')):
    p._element.remove(r)
print("After removing runs, p XML has drawing:", "<w:drawing" in p._element.xml)

run = p.add_run()
run.add_picture(test_img, width=Inches(6.0))
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
print("After adding new picture, p XML has drawing:", "<w:drawing" in p._element.xml)

output_docx = os.path.join(test_dir, "test_replace_out.docx")
doc.save(output_docx)

# Verify with Word COM
import subprocess
ps_script = f"""
$w = New-Object -ComObject Word.Application
$w.Visible = $false
try {{
    $doc = $w.Documents.Open('{output_docx}', $false, $true)
    Write-Host 'Word COM Opened Successfully!'
    Write-Host 'InlineShapes count:' $doc.InlineShapes.Count
    $doc.Close($false)
    Write-Host 'VERIFICATION PASSED: REPLACEMENT 0 XML CORRUPTION!'
}} catch {{
    Write-Host 'VERIFICATION FAILED:' $_.Exception.Message
}} finally {{
    $w.Quit()
}}
"""
ps_file = os.path.join(test_dir, "verify_replace.ps1")
with open(ps_file, "w", encoding="utf-8") as f:
    f.write(ps_script)

res = subprocess.run(["powershell", "-ExecutionPolicy", "Bypass", "-File", ps_file], capture_output=True, text=True)
print("Word COM Output:\n" + res.stdout)
if res.stderr:
    print("Word COM Stderr:\n" + res.stderr)
