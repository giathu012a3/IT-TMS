import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

for fpath in ["c:/IT_System/BaoCao_ITTMS_TKPMHDT (4).docx", "c:/IT_System/BaoCao_ITTMS_TKPMHDT_DotNet9.docx"]:
    print(f"=== {fpath} ===")
    doc = docx.Document(fpath)
    tbl = doc.tables[-1]
    print(f"Rows: {len(tbl.rows)}, Cols: {len(tbl.columns)}")
    for r_idx, r in enumerate(tbl.rows):
        print(f"Row {r_idx}:")
        for c_idx, c in enumerate(r.cells):
            print(f"  Col {c_idx} (width={c.width}):")
            for p in c.paragraphs:
                p_runs = [(run.text, run.bold, run.italic, run.font.name, run.font.size) for run in p.runs]
                print(f"    P style={p.style.name}, align={p.alignment}, runs={p_runs}")
