import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

DOCS = [
    ("Doc1", r"c:\IT_System\BaoCao_ITTMS_TKPMHDT (4).docx"),
    ("Doc2", r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx")
]

OFF_LIMIT_WORDS = ["unit test", "unit-test", "devops", "kiểm thử tự động"]

def verify_doc(label, path):
    print(f"\n==========================================")
    print(f"XÁC MINH PHÁP CHỨNG TÍNH TOÀN VẸN: {label} ({path})")
    print(f"==========================================")
    doc = docx.Document(path)
    
    # 1. Đoạn văn và Tiêu đề
    paras = doc.paragraphs
    headings = [p for p in paras if 'Heading' in p.style.name]
    h1 = [p for p in headings if p.style.name == 'Heading 1']
    h2 = [p for p in headings if p.style.name == 'Heading 2']
    h3 = [p for p in headings if p.style.name == 'Heading 3']
    text_headings = [p for p in headings if p.text.strip()]
    draw_headings = [p for p in headings if not p.text.strip()]
    
    print(f"1. CẤU TRÚC ĐOẠN VĂN & TIÊU ĐỀ:")
    print(f"   - Tổng số Paragraphs: {len(paras)} (Kỳ vọng: 769)")
    assert len(paras) == 769, f"LỖI: Số paragraphs {len(paras)} != 769"
    
    print(f"   - Tổng số Headings: {len(headings)} (Kỳ vọng: 149)")
    print(f"     + Heading 1: {len(h1)} (Kỳ vọng: 14)")
    print(f"     + Heading 2: {len(h2)} (Kỳ vọng: 62)")
    print(f"     + Heading 3: {len(h3)} (Kỳ vọng: 73)")
    print(f"     + Tiêu đề có text: {len(text_headings)} (Kỳ vọng: 136)")
    print(f"     + Heading rỗng (Image holder): {len(draw_headings)} (Kỳ vọng: 13)")
    assert len(headings) == 149, f"LỖI: Headings {len(headings)} != 149"
    assert len(text_headings) == 136, f"LỖI: Text Headings {len(text_headings)} != 136"
    assert len(draw_headings) == 13, f"LỖI: Draw Headings {len(draw_headings)} != 13"
    
    # 2. Cấu trúc Mục lục (TOC)
    toc_paras = [p.text.strip() for p in paras[60:196] if p.text.strip()]
    print(f"2. CẤU TRÚC MỤC LỤC (TOC):")
    print(f"   - Số mục TOC: {len(toc_paras)} (Kỳ vọng: 136)")
    assert len(toc_paras) == 136, f"LỖI: TOC items {len(toc_paras)} != 136"
    
    # 3. Cấu trúc Bảng biểu
    print(f"3. CẤU TRÚC BẢNG BIỂU:")
    print(f"   - Tổng số bảng: {len(doc.tables)} (Kỳ vọng: 55)")
    assert len(doc.tables) == 55, f"LỖI: Số bảng {len(doc.tables)} != 55"
    
    # 4. Kiểm tra Bảng PL.1
    pl1 = doc.tables[-1]
    print(f"4. KIỂM TRA BẢNG PHÂN CÔNG PHỤ LỤC 1 (BẢNG PL.1):")
    print(f"   - Kích thước: {len(pl1.rows)} dòng x {len(pl1.columns)} cột (Kỳ vọng: 5 dòng x 3 cột)")
    assert len(pl1.rows) == 5, f"LỖI: Số dòng {len(pl1.rows)} != 5"
    assert len(pl1.columns) == 3, f"LỖI: Số cột {len(pl1.columns)} != 3"
    
    header = [c.text.strip().replace('\n', ' ') for c in pl1.rows[0].cells]
    print(f"   - Header: {header}")
    assert header[0] == "Thành viên"
    assert header[1] == "MSSV"
    assert "công việc" in header[2].lower()
    
    expected_members = [
        ("Nguyễn Gia Thụ", "25810041", "Staff", "StaffController.cs", "State Pattern", "5.2", "100%"),
        ("Nguyễn Mạnh Tiến", "25810043", "Leader", "LeaderController.cs", "Factory", "5.1", "100%"),
        ("Hồ Thanh Tình", "25810044", "User", "TicketController.cs", "Observer", "TicketHub", "100%"),
        ("Trần Thanh Danh", "25810009", "Admin", "AdminController.cs", "Decorator", "ERD", "100%")
    ]
    
    all_table_text = ""
    for r_idx, exp in enumerate(expected_members, start=1):
        row = pl1.rows[r_idx]
        name = row.cells[0].text.strip()
        mssv = row.cells[1].text.strip()
        work = row.cells[2].text.strip()
        all_table_text += " " + name + " " + mssv + " " + work
        
        print(f"   - Dòng {r_idx}: {name} ({mssv})")
        assert exp[0] in name, f"LỖI: Tên {name} không chứa {exp[0]}"
        assert exp[1] in mssv, f"LỖI: MSSV {mssv} không chứa {exp[1]}"
        for kw in exp[2:]:
            assert kw.lower() in work.lower(), f"LỖI: Dòng của {name} thiếu từ khóa {kw}"
        print(f"     + Đầy đủ Vertical Slice: {exp[2]}, Controller, Mẫu thiết kế, Sơ đồ, 100% đóng góp")
        
    # Kiểm tra loại bỏ triệt để từ khóa ngoài lề
    print(f"   - Kiểm tra loại bỏ triệt để devops/unit test:")
    for bad_word in OFF_LIMIT_WORDS:
        assert bad_word not in all_table_text.lower(), f"LỖI: Bảng PL.1 còn chứa từ khóa cấm: '{bad_word}'"
    print(f"     + XÁC NHẬN: 0 từ khóa ngoài lề ({', '.join(OFF_LIMIT_WORDS)}) tồn tại trong Bảng PL.1!")
    
    # 5. Kiểm tra không còn bảng tạm thừa [CHÈN HÌNH] nào
    for idx, tbl in enumerate(doc.tables):
        tbl_text = "".join(tbl._tbl.itertext()).strip()
        assert "[CHÈN HÌNH]" not in tbl_text, f"LỖI: Còn bảng tạm [CHÈN HÌNH] tại table index {idx}"
    print(f"5. XÁC NHẬN BẢNG TẠM: 0 bảng tạm thừa [CHÈN HÌNH] tồn tại trong tài liệu!")
    print(f"[VERDICT] {label}: TOÀN BỘ CÁC TIÊU CHÍ ĐỀU ĐẠT CHUẨN 100% (PASS)")

if __name__ == "__main__":
    for label, path in DOCS:
        verify_doc(label, path)
    print("\n>>> TẤT CẢ CÁC BÀI KIỂM THỬ XÁC MINH TRÊN PYTHON ĐỀU THÀNH CÔNG RỰC RỠ! <<<")
