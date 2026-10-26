import sys
import docx

sys.stdout.reconfigure(encoding='utf-8')

def compare_body_elements():
    doc1 = docx.Document("c:/IT_System/BaoCao_ITTMS_TKPMHDT (4).docx")
    doc2 = docx.Document("c:/IT_System/BaoCao_ITTMS_TKPMHDT_DotNet9.docx")
    
    body1 = doc1._body._element
    body2 = doc2._body._element
    
    print(f"Doc1 body child count: {len(body1)}")
    print(f"Doc2 body child count: {len(body2)}")
    
    # Check elements around Table 42 in Doc2
    tbl42 = doc2.tables[42]._tbl
    idx2 = body2.index(tbl42)
    print(f"\nDoc2 elements around Table 42 (index {idx2}):")
    for i in range(max(0, idx2 - 3), min(len(body2), idx2 + 4)):
        el = body2[i]
        tag = el.tag.split('}')[-1]
        text = el.text[:50] if el.text else ""
        if tag == 'p':
            p_text = "".join(el.itertext()).strip()
            print(f"  [{i}] <p> text: {p_text[:60]}")
        elif tag == 'tbl':
            tbl_text = "".join(el.itertext()).strip()
            print(f"  [{i}] <tbl> text: {tbl_text[:60]}")
        else:
            print(f"  [{i}] <{tag}>")

compare_body_elements()
