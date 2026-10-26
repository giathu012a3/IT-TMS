import sys
import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Pt

sys.stdout.reconfigure(encoding='utf-8')

DOC1_PATH = r"c:\IT_System\BaoCao_ITTMS_TKPMHDT (4).docx"
DOC2_PATH = r"c:\IT_System\BaoCao_ITTMS_TKPMHDT_DotNet9.docx"

COL_WIDTHS = [2200, 1300, 5526] # dxa, total = 9026

MEMBERS_DATA = [
    {
        "name": "Nguyễn Gia Thụ",
        "mssv": "25810041",
        "role": "Kỹ thuật viên (Staff)",
        "code": "StaffController.cs, Views/Staff/ (Dashboard, Assigned, Details, History)",
        "pattern": "State Pattern (quản lý 7 trạng thái vòng đời ticket: New, Assigned, InProgress, Waiting, Resolved, Closed, Rejected; đóng băng các thao tác sai nghiệp vụ qua Result<T>)",
        "diagrams": "Sơ đồ chuyển trạng thái (State Machine – Hình 5.2), Lược đồ lớp State Pattern (Hình 6.1), Lược đồ tuần tự UC07 – Cập nhật trạng thái xử lý (Hình 5.7)",
        "contrib": "100% (Hoàn thành xuất sắc nhiệm vụ)"
    },
    {
        "name": "Nguyễn Mạnh Tiến",
        "mssv": "25810043",
        "role": "Trưởng nhóm (Leader)",
        "code": "LeaderController.cs, Views/Leader/ (Dashboard, NewQueue, Index, Details)",
        "pattern": "Factory Method / Flyweight Pattern (TicketStateFactory.cs – khởi tạo và tái sử dụng đối tượng trạng thái, tối ưu bộ nhớ)",
        "diagrams": "Quy trình xử lý ticket hoàn thiện (BPMN – Hình 5.1), Lược đồ tuần tự UC09 – Phân công yêu cầu (Hình 5.6), Lược đồ tuần tự UC10 – Từ chối yêu cầu (Hình 5.9)",
        "contrib": "100% (Hoàn thành xuất sắc nhiệm vụ)"
    },
    {
        "name": "Hồ Thanh Tình",
        "mssv": "25810044",
        "role": "Người dùng (User)",
        "code": "TicketController.cs, Views/Ticket/ (Dashboard, Index, Create, Details)",
        "pattern": "Observer Pattern (phát Domain Events, đẩy thông báo thời gian thực qua SignalR WebSocket TicketHub và tự động ghi nhật ký kiểm toán SystemLog)",
        "diagrams": "Lược đồ tuần tự UC03 – Tạo yêu cầu mới (Hình 5.5), Lược đồ tuần tự UC06 – Trao đổi thông tin ticket (Hình 5.8), Lược đồ tuần tự UC19, UC04 – Đóng/Mở lại yêu cầu và đánh giá (Hình 5.10)",
        "contrib": "100% (Hoàn thành xuất sắc nhiệm vụ)"
    },
    {
        "name": "Trần Thanh Danh",
        "mssv": "25810009",
        "role": "Quản trị & Bảo mật (Admin & Auth)",
        "code": "AdminController.cs, AuthController.cs, Views/Admin/, Views/Auth/",
        "pattern": "Decorator Pattern ([AuthorizeRoles] kiểm soát truy cập RBAC và chặn tài khoản bị vô hiệu hóa), Singleton Pattern (PasswordHasherService băm và xác thực mật khẩu PBKDF2 HMAC-SHA512)",
        "diagrams": "Lược đồ thực thể – quan hệ ERD 8 bảng CSDL (Hình 3.1), Lược đồ lớp tổng quát hệ thống (Class Diagram – Hình 5.3), Toàn bộ 5 Lược đồ Use Case tổng thể và chi tiết (Hình 2.1 – Hình 2.5)",
        "contrib": "100% (Hoàn thành xuất sắc nhiệm vụ)"
    }
]

def clean_doc2_placeholder_tables(doc):
    """Xóa bỏ 7 bảng tạm [CHÈN HÌNH] thừa trong Doc2"""
    to_delete = []
    for idx, tbl in enumerate(doc.tables):
        text = "".join(tbl._tbl.itertext()).strip()
        if "[CHÈN HÌNH]" in text:
            to_delete.append((idx, tbl))
            
    print(f"[*] Doc2: Tìm thấy {len(to_delete)} bảng tạm [CHÈN HÌNH] để dọn dẹp.")
    for idx, tbl in to_delete:
        parent = tbl._tbl.getparent()
        parent.remove(tbl._tbl)
    print(f"[+] Doc2: Đã xóa thành công {len(to_delete)} bảng tạm thừa.")

def update_table_pl1(doc, doc_name=""):
    """Chuẩn hóa Bảng PL.1 theo ma trận 4 Vertical Slices độc lập"""
    pl1 = doc.tables[-1]
    print(f"[*] Cập nhật Bảng PL.1 cho {doc_name} (Hiện có {len(pl1.rows)} dòng, {len(pl1.columns)} cột)...")
    
    # Header row
    header_titles = ["Thành viên", "MSSV", "Nội dung công việc đảm nhiệm"]
    for c_idx, title in enumerate(header_titles):
        cell = pl1.rows[0].cells[c_idx]
        cell.text = ""
        p = cell.paragraphs[0]
        p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(title)
        r.bold = True
        r.font.size = Pt(11.5)
        
    # Data rows
    for r_idx, data in enumerate(MEMBERS_DATA):
        row = pl1.rows[r_idx + 1]
        
        # Col 0: Tên thành viên
        c0 = row.cells[0]
        c0.text = ""
        p0 = c0.paragraphs[0]
        p0.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
        p0.paragraph_format.space_before = Pt(0)
        p0.paragraph_format.space_after = Pt(2)
        r0 = p0.add_run(data["name"])
        r0.bold = True
        r0.font.size = Pt(11)
        
        # Col 1: MSSV
        c1 = row.cells[1]
        c1.text = ""
        p1 = c1.paragraphs[0]
        p1.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.CENTER
        p1.paragraph_format.space_before = Pt(0)
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(data["mssv"])
        r1.bold = False
        r1.font.size = Pt(11)
        
        # Col 2: Nội dung công việc theo Vertical Slice
        c2 = row.cells[2]
        c2.text = ""
        work_items = [
            ("• Phân hệ: ", data["role"]),
            ("• Mã nguồn sở hữu: ", data["code"]),
            ("• Mẫu thiết kế GoF: ", data["pattern"]),
            ("• Sơ đồ kỹ thuật: ", data["diagrams"]),
            ("• Tỷ lệ đóng góp: ", data["contrib"])
        ]
        for item_idx, (prefix, content) in enumerate(work_items):
            p = c2.paragraphs[0] if item_idx == 0 else c2.add_paragraph()
            p.alignment = docx.enum.text.WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            r_pre = p.add_run(prefix)
            r_pre.bold = True
            r_pre.font.size = Pt(11)
            
            r_cnt = p.add_run(content)
            r_cnt.bold = False
            r_cnt.font.size = Pt(11)
            
    # Chuẩn hóa độ rộng cột và padding cho từng cell
    for row in pl1.rows:
        for c_idx, w in enumerate(COL_WIDTHS):
            tc = row.cells[c_idx]
            tcPr = tc._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn('w:tcW'))
            if tcW is not None:
                tcW.set(qn('w:w'), str(w))
            else:
                tcPr.append(parse_xml(f'<w:tcW {nsdecls("w")} w:w="{w}" w:type="dxa"/>'))
                
            # Đảm bảo padding chuẩn
            tcMar = tcPr.find(qn('w:tcMar'))
            if tcMar is None:
                tcPr.append(parse_xml(
                    f'<w:tcMar {nsdecls("w")}>'
                    f'<w:top w:w="50" w:type="dxa"/>'
                    f'<w:left w:w="90" w:type="dxa"/>'
                    f'<w:bottom w:w="50" w:type="dxa"/>'
                    f'<w:right w:w="90" w:type="dxa"/>'
                    f'</w:tcMar>'
                ))
    print(f"[+] {doc_name}: Chuẩn hóa Bảng PL.1 hoàn tất thành công!")

def process_file(path, is_doc2=False):
    doc_name = "Doc2 (DotNet9)" if is_doc2 else "Doc1 (BaoCao 4)"
    print(f"\n==========================================")
    print(f"Bắt đầu xử lý {doc_name}: {path}")
    doc = docx.Document(path)
    
    if is_doc2:
        clean_doc2_placeholder_tables(doc)
        
    update_table_pl1(doc, doc_name)
    
    doc.save(path)
    print(f"[SUCCESS] Đã lưu file: {path}")

if __name__ == "__main__":
    process_file(DOC1_PATH, is_doc2=False)
    process_file(DOC2_PATH, is_doc2=True)
    print("\n[COMPLETE] Đã hoàn thành cập nhật cả 2 tài liệu Word!")
