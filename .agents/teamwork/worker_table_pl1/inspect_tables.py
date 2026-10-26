import sys
import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

sys.stdout.reconfigure(encoding='utf-8')

def inspect_doc(path):
    print(f"=== Inspecting {path} ===")
    doc = docx.Document(path)
    print(f"Total tables: {len(doc.tables)}")
    
    # Check last table (Table PL.1)
    pl1 = doc.tables[-1]
    print(f"Last table ({len(doc.tables)-1}) rows: {len(pl1.rows)}, cols: {len(pl1.columns)}")
    for r_idx, row in enumerate(pl1.rows):
        cells_text = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"Row {r_idx}: {cells_text}")
    
    # Inspect XML properties of PL.1
    tblPr = pl1._tbl.tblPr
    print("tblPr xml:", tblPr.xml[:300] if tblPr is not None else "None")
    
    # If Doc2, find the placeholder tables
    if "DotNet9" in path:
        print("\n--- Scanning for [CHÈN HÌNH] tables in Doc2 ---")
        for idx, tbl in enumerate(doc.tables):
            first_cell = tbl.rows[0].cells[0].text.strip() if tbl.rows and tbl.rows[0].cells else ""
            if "[CHÈN HÌNH]" in first_cell or "CHÈN HÌNH" in first_cell:
                print(f"Table index {idx}: {first_cell[:80]}...")
                # find parent element and tag
                parent = tbl._tbl.getparent()
                print(f"  Parent tag: {parent.tag}, index in parent: {parent.index(tbl._tbl)}")

if __name__ == "__main__":
    inspect_doc("c:/IT_System/BaoCao_ITTMS_TKPMHDT (4).docx")
    inspect_doc("c:/IT_System/BaoCao_ITTMS_TKPMHDT_DotNet9.docx")
