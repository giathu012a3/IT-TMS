import sys
import shutil
import docx
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')

# Copy to test file
test_doc2 = "c:/IT_System/test_doc2.docx"
shutil.copy2("c:/IT_System/BaoCao_ITTMS_TKPMHDT_DotNet9.docx", test_doc2)

doc = docx.Document(test_doc2)
print("Initial tables in test_doc2:", len(doc.tables))
print("Initial paragraphs in test_doc2:", len(doc.paragraphs))

# Identify the 7 [CHÈN HÌNH] tables
to_delete = []
for idx, tbl in enumerate(doc.tables):
    text = "".join(tbl._tbl.itertext()).strip()
    if "[CHÈN HÌNH]" in text:
        to_delete.append((idx, tbl))

print(f"Found {len(to_delete)} placeholder tables to delete.")
for idx, tbl in to_delete:
    parent = tbl._tbl.getparent()
    parent.remove(tbl._tbl)

doc.save(test_doc2)

# Reload and verify
doc_reloaded = docx.Document(test_doc2)
print("After deletion tables:", len(doc_reloaded.tables))
print("After deletion paragraphs:", len(doc_reloaded.paragraphs))

headings = [p for p in doc_reloaded.paragraphs if 'Heading' in p.style.name]
text_headings = [p for p in headings if p.text.strip()]
print(f"Headings count: {len(headings)} (Text={len(text_headings)})")
